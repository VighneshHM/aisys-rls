"""Read .csv / .xlsx into (headers, rows-of-strings). Everything becomes a trimmed string; blanks are ''."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import PurePath

from openpyxl import load_workbook

MAX_BYTES = 30 * 1024 * 1024
MAX_ROWS = 200_000
MAX_COLUMNS = 100


class ParseError(ValueError):
    """The file cannot be read as a table (user-facing message)."""


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 2019 as 2019.0 and long IDs as floats
    return str(value).strip()


def _unique_headers(raw: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(raw, start=1):
        name = h.strip() or f"column_{i}"
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return out


def _from_csv(data: bytes) -> tuple[list[str], list[tuple[int, list[str]]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        rows = [[c.strip() for c in r] for r in reader]
    except csv.Error as exc:  # e.g. one cell over the csv module's 131,072-character limit (D-072)
        raise ParseError(f"The CSV file cannot be read near line {reader.line_num}: {exc}") from exc
    return _split_header(rows)


def _from_xlsx(data: bytes) -> tuple[list[str], list[tuple[int, list[str]]]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise ParseError("The file is not a valid .xlsx workbook") from exc
    try:
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            raise ParseError("The workbook has no sheets")
        rows: list[list[str]] = []
        for r in ws.iter_rows(values_only=True):
            rows.append([_cell(c) for c in r])
            if len(rows) > MAX_ROWS + 1:
                raise ParseError(f"Too many rows (limit {MAX_ROWS})")
    finally:
        wb.close()
    return _split_header(rows)


def _split_header(rows: list[list[str]]) -> tuple[list[str], list[tuple[int, list[str]]]]:
    """Return headers and (spreadsheet row number, cells); numbers match what Excel shows."""
    numbered = list(enumerate(rows, start=1))
    while numbered and not any(c for c in numbered[0][1]):
        numbered.pop(0)  # leading blank rows
    if not numbered:
        raise ParseError("The file is empty")
    header_row = numbered[0][1]
    width = len(header_row)
    while width and not header_row[width - 1]:
        width -= 1  # ignore trailing empty header cells
    if width == 0:
        raise ParseError("No header row found")
    if width > MAX_COLUMNS:
        raise ParseError(f"Too many columns (limit {MAX_COLUMNS})")
    headers = _unique_headers(header_row[:width])
    body = []
    for no, r in numbered[1:]:
        cells = (r + [""] * width)[:width]
        if any(cells):  # skip fully blank lines
            body.append((no, cells))
    return headers, body


def read_table(filename: str, data: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    """Return (headers, [(spreadsheet_row_number, {header: value})])."""
    if len(data) > MAX_BYTES:
        raise ParseError(f"File too large (limit {MAX_BYTES // (1024 * 1024)} MB)")
    ext = PurePath(filename).suffix.lower()
    if ext == ".csv":
        headers, body = _from_csv(data)
    elif ext == ".xlsx":
        headers, body = _from_xlsx(data)
    else:
        raise ParseError("Unsupported file type: use .xlsx or .csv")
    if len(body) > MAX_ROWS:
        raise ParseError(f"Too many rows (limit {MAX_ROWS})")
    return headers, [(no, dict(zip(headers, cells, strict=True))) for no, cells in body]
